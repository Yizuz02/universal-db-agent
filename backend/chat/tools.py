import json
import os
import sqlite3
from pydantic import BaseModel, Field
from langchain_core.tools import tool
from datetime import datetime
import csv
from pathlib import Path

# --- PYDANTIC SCHEMAS FOR TOOLS ---

class QueryInput(BaseModel):
    query: str = Field(
        description="The exact SQL SELECT query to execute. Must be valid SQLite syntax and read-only."
    )

class SaveReportInput(BaseModel):
    filename: str = Field(
        description="The name of the file to save (e.g., 'project_summary.md'). Must include the .md extension."
    )
    content: str = Field(
        description="The complete and well-formatted Markdown string containing the report data, titles, and tables."
    )

class SaveCSVReportInput(BaseModel):
    query: str = Field(
        description="The exact SQL SELECT query to execute to fetch the data for the CSV file. Must be read-only."
    )
    filename: str = Field(
        description="The name of the CSV file to generate (e.g., 'work_hours_report.csv'). Must include the .csv extension."
    )

class SaveExcelReportInput(BaseModel):
    query: str = Field(
        description="The exact SQL SELECT query to execute to fetch the data for the Excel file. Must be read-only."
    )
    filename: str = Field(
        description="The name of the Excel file to generate (e.g., 'project_report.xlsx'). Must include the .xlsx extension."
    )

# --- TOOLS DEFINITIONS ---

def build_tools(db_path: Path, reports_dir: Path, metadata_path: Path, on_file_saved=None):
    @tool
    def list_tables_and_columns() -> str:
        """
        Fetches and inspects the database to retrieve the names of ALL existing tables, 
        their respective columns, and data types. It also injects business definitions 
        from the metadata dictionary. Always use this tool first if you do not know 
        the database schema structure.
        """
        try:
            # Load external metadata if available
            metadata = {}
            if metadata_path.exists():
                with open(metadata_path, "r", encoding="utf-8") as f:
                    metadata = json.load(f)
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';")
            tables = [row[0] for row in cursor.fetchall()]
            
            schema_info = "Database Schema Structure & Business Dictionary:\n"
            if "database_description" in metadata:
                schema_info += f"Context: {metadata['database_description']}\n\n"
                
            for table in tables:
                cursor.execute(f"PRAGMA table_info({table});")
                columns = [f"{col[1]} ({col[2]})" for col in cursor.fetchall()]
                
                schema_info += f"- Table '{table}': Columns -> {', '.join(columns)}\n"
                
                # Inject business descriptions from JSON
                if "tables" in metadata and table in metadata["tables"]:
                    t_meta = metadata["tables"][table]
                    if "description" in t_meta:
                        schema_info += f"  Description: {t_meta['description']}\n"
                    if "column_annotations" in t_meta:
                        for col, annot in t_meta["column_annotations"].items():
                            schema_info += f"  Annotation [{col}]: {annot}\n"
                            
            conn.close()
            return schema_info
        except Exception as e:
            return f"Error reading database structure or metadata: {str(e)}"


    @tool(args_schema=QueryInput)
    def execute_read_only_query(query: str) -> str:
        """
        Executes a SQL SELECT query against the database and returns the results.
        CRITICAL: Only read-only queries (SELECT) are allowed. Do not attempt to modify or delete data.
        """
        query_clean = query.strip().lower()
        if not query_clean.startswith("select"):
            return "Error: For security reasons, only read-only queries (SELECT) are permitted."
            
        try:
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            cursor.execute(query)
            
            columns = [desc[0] for desc in cursor.description]
            rows = cursor.fetchall()
            
            conn.close()
            
            if not rows:
                return "The query executed successfully but returned no records."
                
            result = f"Query Results (Columns: {', '.join(columns)}):\n"
            for row in rows:
                result += f"{str(row)}\n"
            return result
            
        except sqlite3.Error as e:
            # We return the exact database error so the ReAct loop can attempt to self-correct
            return f"SQL Execution Error: {str(e)}. Please review the table names/columns and try again."
        
    @tool(args_schema=SaveReportInput)
    def save_markdown_report(filename: str, content: str) -> str:
        """
        Creates and saves a structured Markdown (.md) report file into the local reports directory.
        Use this tool ONLY when the user explicitly requests to save, export, or generate a file/report.
        Do not use this tool if the user just wants to see the answer in the chat text.
        """
        file_extension = Path(filename).suffix
        if file_extension.lower() != ".md":
            return f"Error: Only Markdown files (.md) are supported for reports. Received: {file_extension}"

        target_dir = reports_dir / "markdown"
        
        try:
            target_dir.mkdir(parents=True, exist_ok=True)
                
            # Secure the path to prevent directory traversal attacks
            safe_filename = Path(filename).name
            file_path = target_dir / safe_filename
            
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(content)
                
            if on_file_saved:
                on_file_saved(str(file_path))
                
            return f"Success: Report successfully saved to '{file_path}'."
        except Exception as e:
            return f"Error saving the report file: {str(e)}"
        

    @tool
    def get_current_datetime() -> str:
        """
        Returns the current date and time of the system. 
        Use this tool whenever you need to timestamp reports, check the current year, 
        or calculate time differences.
        """
        try:
            now = datetime.now()
            # Format example: "Saturday, July 18, 2026, 13:34 PM"
            return f"Current Date and Time: {now.strftime('%A, %B %d, %Y, %H:%M %p')}"
        except Exception as e:
            return f"Error fetching current datetime: {str(e)}"
        
    @tool(args_schema=SaveCSVReportInput)
    def save_csv_report(query: str, filename: str) -> str:
        """
        Executes a SQL SELECT query, extracts the raw records, and exports them directly 
        into a clean CSV spreadsheet file inside the local exports directory.
        Use this tool ONLY when the user explicitly asks to export data to a CSV file.
        """
        query_clean = query.strip().lower()
        if not query_clean.startswith("select"):
            return "Error: For security reasons, only read-only queries (SELECT) are permitted for CSV exports."
            
        file_extension = Path(filename).suffix
        if file_extension.lower() != ".csv":
            return f"Error: Only CSV files (.csv) are supported for CSV exports. Received: {file_extension}"

        target_dir = reports_dir / "csv"

        try:
            target_dir.mkdir(parents=True, exist_ok=True)
                
            safe_filename = Path(filename).name
            file_path = target_dir / safe_filename
            
            # Connect to DB and execute
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            cursor.execute(query)
            
            headers = [desc[0] for desc in cursor.description]
            rows = cursor.fetchall()
            conn.close()
            
            if not rows:
                return "The query executed successfully but returned no data to export."
                
            # Write to CSV file
            with open(file_path, mode="w", newline="", encoding="utf-8") as csv_file:
                writer = csv.writer(csv_file)
                writer.writerow(headers) # Write headers first
                writer.writerows(rows)   # Write data rows
                
            if on_file_saved:
                on_file_saved(str(file_path))
                
            return f"Success: Data successfully exported to a clean CSV sheet at '{file_path}' ({len(rows)} records exported)."
            
        except sqlite3.Error as e:
            return f"SQL Execution Error during CSV export: {str(e)}. Please check syntax and try again."
        except Exception as e:
            return f"Error creating CSV file: {str(e)}"

    @tool(args_schema=SaveExcelReportInput)
    def save_excel_report(query: str, filename: str) -> str:
        """
        Executes a SQL SELECT query, extracts the records, and exports them into a real Excel (.xlsx) spreadsheet file.
        Use this tool ONLY when the user explicitly asks to save to an Excel file.
        """
        import openpyxl
        from openpyxl.styles import Font
        
        query_clean = query.strip().lower()
        if not query_clean.startswith("select"):
            return "Error: For security reasons, only read-only queries (SELECT) are permitted for Excel exports."
            
        file_extension = Path(filename).suffix
        if file_extension.lower() != ".xlsx":
            return f"Error: Only Excel files (.xlsx) are supported for Excel exports. Received: {file_extension}"

        target_dir = reports_dir / "excel"

        try:
            target_dir.mkdir(parents=True, exist_ok=True)
                
            safe_filename = Path(filename).name
            file_path = target_dir / safe_filename
            
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            cursor.execute(query)
            
            headers = [desc[0] for desc in cursor.description]
            rows = cursor.fetchall()
            conn.close()
            
            if not rows:
                return "The query executed successfully but returned no data to export."
                
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Report"
            
            bold_font = Font(bold=True)
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = bold_font
            
            for row_idx, row in enumerate(rows, 2):
                for col_idx, value in enumerate(row, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
                    
            wb.save(str(file_path))
            
            if on_file_saved:
                on_file_saved(str(file_path))
                
            return f"Success: Data successfully exported to an Excel file at '{file_path}' ({len(rows)} records exported)."
            
        except sqlite3.Error as e:
            return f"SQL Execution Error during Excel export: {str(e)}. Please check syntax and try again."
        except Exception as e:
            return f"Error creating Excel file: {str(e)}"

    return [list_tables_and_columns, execute_read_only_query, save_markdown_report, get_current_datetime, save_csv_report, save_excel_report]
